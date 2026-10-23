import ast
from pathlib import Path
from html.parser import HTMLParser
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
class Content(HTMLParser):
    def __init__(self,p): super().__init__(); self.p=p; self.bold=False
    def handle_starttag(self,t,a):
        if t=='b': self.bold=True
        elif t=='br': self.p.add_run().add_break()
    def handle_endtag(self,t):
        if t=='b': self.bold=False
    def handle_data(self,d): self.p.add_run(d).bold=self.bold

def fill(p,t): Content(p).feed(t)
d=Document()
sec=d.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55)
sec.left_margin=Inches(.61); sec.right_margin=Inches(.61)
for name in ['Normal','Title','Heading 1','Heading 2']:
    st=d.styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(9.5)
d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.05
for name,size in [('Title',15),('Heading 1',15),('Heading 2',10.5)]:
    st=d.styles[name]; st.font.size=Pt(size); st.font.bold=True
    st.paragraph_format.space_before=Pt(8); st.paragraph_format.space_after=Pt(5)
first=True
for node in ast.parse(Path('tmp/pdfs/build_use_case.py').read_text()).body:
    if not isinstance(node,ast.Expr) or not isinstance(node.value,ast.Call): continue
    call=node.value
    if not isinstance(call.func,ast.Name) or call.func.id not in ['page','h','p','table']: continue
    value=ast.literal_eval(call.args[0]); kind=call.func.id
    if kind=='page':
        if not first: d.add_page_break()
        d.add_paragraph(value,'Title' if first else 'Heading 1'); first=False
    elif kind=='h': d.add_paragraph(value,'Heading 2')
    elif kind=='p': fill(d.add_paragraph(),value)
    else:
        t=d.add_table(rows=0,cols=2); t.autofit=False
        t.columns[0].width=Inches(2.01); t.columns[1].width=Inches(5.27)
        for idx,row in enumerate(value):
            cells=t.add_row().cells
            for j,txt in enumerate(row):
                cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                fill(cells[j].paragraphs[0],txt)
                pr=cells[j]._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
                pr.append(borders)
                margins=OxmlElement('w:tcMar')
                for edge in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+edge); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); margins.append(el)
                pr.append(margins)
                if idx==0:
                    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'EEEEEE'); pr.append(sh)
            rowpr=t.rows[-1]._tr.get_or_add_trPr()
            rowpr.append(OxmlElement('w:cantSplit'))
            if idx==0: rowpr.append(OxmlElement('w:tblHeader'))
        d.add_paragraph().paragraph_format.space_after=Pt(0)
sec.footer.paragraphs[0].text='UC-04 | Show useful sponsored products | CS3203 Fall 2026'
sec.footer.paragraphs[0].style=d.styles['Normal']
for r in sec.footer.paragraphs[0].runs: r.font.size=Pt(8)
Path('output/doc').mkdir(parents=True,exist_ok=True)
out=Path('output/doc/StudentID_Assignment1UseCases_CS3203Fall2026.docx')
d.save(out)
print(out.resolve())
